#!/usr/bin/env python3
"""
build_workbook.py
Builds Northfield_FY2025_Workpapers.xlsx from the clean/ tables.

The workbook does the mechanical parts with live formulas (tie-outs, monthly analytics,
lookups to subsequent receipts and the A/P listing). Every judgment is left to you:
expectations, sample selection, whether an item is a liability, conclusions.
Yellow cells with blue text = your inputs.

Usage
  python build_workbook.py --clean clean --out Northfield_FY2025_Workpapers.xlsx
Then open in Excel (it recalculates on open).
"""
import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
FILL_HDR = PatternFill("solid", fgColor="1F3864")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_EXAMPLE = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
MONEY = '#,##0.00;(#,##0.00);"-"'
MONEY0 = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "mm/dd/yyyy"


def style_range(ws, ref, font=None, fill=None, fmt=None, border=None, align=None):
    for row in ws[ref]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if fmt: c.number_format = fmt
            if border: c.border = border
            if align: c.alignment = align


def header_row(ws, row, headers, col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font, c.fill, c.border = F_HDR, FILL_HDR, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")


def put(ws, ref, value, font=F_BASE, fmt=None, fill=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    return c


def inp(ws, ref, value=None, fmt=None):
    return put(ws, ref, value, F_INPUT, fmt, FILL_INPUT)


def signoff_row(ws, row=3):
    put(ws, f"A{row}", "Prepared by / date:")
    inp(ws, f"B{row}")
    put(ws, f"D{row}", "Reviewed by / date:")
    inp(ws, f"E{row}")


def add_list_validation(ws, cell_range, choices):
    dv = DataValidation(type="list", formula1=f'"{",".join(choices)}"', allow_blank=True)
    dv.error = f"Choose one of: {', '.join(choices)}"
    dv.errorTitle = "Invalid selection"
    dv.prompt = f"Choose one of: {', '.join(choices)}"
    dv.promptTitle = "Required judgment"
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(cell_range)


def section(ws, row, text, width=10):
    for col in range(1, width + 1):
        ws.cell(row=row, column=col).fill = FILL_SECTION
    c = ws.cell(row=row, column=1, value=text)
    c.font = F_BOLD


def wp_header(ws, title, ref):
    put(ws, "A1", f"Northfield Supply Co. | FY2025 | {ref}", F_TITLE)
    put(ws, "A2", title, F_BOLD)
    signoff_row(ws)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def write_data(wb, name, df, date_cols=(), money_cols=()):
    ws = wb.create_sheet(name)
    header_row(ws, 1, list(df.columns))
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        for c_idx, v in enumerate(row, start=1):
            if v is pd.NaT or (isinstance(v, float) and pd.isna(v)):
                v = None
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime() if not pd.isna(v) else None
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.font = F_BASE
    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(12, min(40, len(col) + 4))
        if col in date_cols:
            for c in ws[letter][1:]:
                c.number_format = DATE
        if col in money_cols:
            for c in ws[letter][1:]:
                c.number_format = MONEY
    ws.freeze_panes = "A2"
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--clean", default="clean")
    ap.add_argument("--out", default="Northfield_FY2025_Workpapers.xlsx")
    args = ap.parse_args()
    cd = Path(args.clean)

    tb = pd.read_csv(cd / "trial_balance.csv", dtype={"account": str})
    sales = pd.read_csv(cd / "sales_ledger.csv", parse_dates=["inv_date", "ship_date"])
    ship = pd.read_csv(cd / "shipping_documents.csv", parse_dates=["ship_date"])
    aging = pd.read_csv(cd / "ar_aging.csv", parse_dates=["inv_date", "due_date"])
    rec = pd.read_csv(cd / "cash_receipts.csv", parse_dates=["receipt_date"])
    apl = pd.read_csv(cd / "ap_listing.csv", parse_dates=["inv_date", "due_date"])
    jd = pd.read_csv(cd / "jan_disbursements.csv", parse_dates=["paid_date", "inv_date"], dtype={"check_no": str})
    pym = pd.read_csv(cd / "prior_year_monthly.csv")

    wb = Workbook()
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    idx = wb.active
    idx.title = "Index"
    for n in ["Planning", "Lead", "WP1_Revenue", "WP2_AR", "WP3_SURL", "Findings",
              "Review_Notes", "AI_Log", "Time_Log"]:
        wb.create_sheet(n)
    write_data(wb, "D_TB", tb, money_cols=("cy_balance", "py_balance"))
    write_data(wb, "D_Sales", sales, date_cols=("inv_date", "ship_date"), money_cols=("amount",))
    write_data(wb, "D_Shipping", ship, date_cols=("ship_date",))
    write_data(wb, "D_AR", aging, date_cols=("inv_date", "due_date"),
               money_cols=("balance", "b_0_30", "b_31_60", "b_61_90", "b_91_120", "b_over_120"))
    write_data(wb, "D_Receipts", rec, date_cols=("receipt_date",), money_cols=("amount",))
    write_data(wb, "D_AP", apl, date_cols=("inv_date", "due_date"), money_cols=("amount",))
    write_data(wb, "D_JanDisb", jd, date_cols=("paid_date", "inv_date"), money_cols=("amount",))
    write_data(wb, "D_PYMonthly", pym, money_cols=("py_net_sales",))

    # ================================================================ Index
    ws = idx
    put(ws, "A1", "Northfield Supply Co. | FY2025 Financial Statement Audit (mini)", F_TITLE)
    put(ws, "A2", "FICTIONAL client, synthetic data. Built for an AI-assisted audit portfolio project.", F_NOTE)
    header_row(ws, 4, ["Sheet", "Purpose", "Standard"])
    index_rows = [
        ("Planning", "Materiality and significant risks", "AU-C 320, 315, 240"),
        ("Lead", "Lead sheet: TB mapped to financial statement lines, CY vs PY", "AU-C 520 (planning analytics)"),
        ("WP1_Revenue", "Monthly substantive analytics, cutoff, duplicate test", "AU-C 520, 330, 240"),
        ("WP2_AR", "A/R tie-out, sample + subsequent receipts, allowance", "AU-C 530, 500, 540"),
        ("WP3_SURL", "Search for unrecorded liabilities (January disbursements)", "AU-C 330, 560"),
        ("Findings", "Misstatement log (feeds scoring and the summary of misstatements)", "AU-C 450"),
        ("Review_Notes", "Review comments from the reviewer agent and self-review", "AU-C 230"),
        ("AI_Log", "Where the AI was wrong, how you caught it, product feedback", "-"),
        ("Time_Log", "Planned vs actual minutes per block", "-"),
        ("D_* sheets", "Clean client data from normalize.py. Do not edit.", "-"),
    ]
    for i, r in enumerate(index_rows, start=5):
        for j, v in enumerate(r, start=1):
            ws.cell(row=i, column=j, value=v).font = F_BASE
    section(ws, 16, "Legend", 3)
    put(ws, "A17", "Your input", F_INPUT, fill=FILL_INPUT); put(ws, "B17", "Yellow cell, blue text: edit these only")
    put(ws, "A18", "Formula", F_BASE); put(ws, "B18", "Black text: calculated, do not overwrite")
    put(ws, "A19", "Cross-sheet link", F_LINK); put(ws, "B19", "Green text: pulls from another sheet")
    section(ws, 21, "Tickmark legend", 3)
    ticks = [("a", "Agreed to source document"), ("F", "Footed"), ("CF", "Cross-footed"),
             ("T", "Traced to trial balance / lead sheet"), ("S", "Agreed to subsequent cash receipt"),
             ("SD", "Agreed to shipping document"), ("X", "Exception, see Findings")]
    for i, (t, d) in enumerate(ticks, start=22):
        put(ws, f"A{i}", t, F_BOLD); put(ws, f"B{i}", d)
    widths(ws, {"A": 22, "B": 70, "C": 30})

    # ================================================================ Lead sheet
    ws = wb["Lead"]
    put(ws, "A1", "Northfield Supply Co. | FY2025 | Lead Sheet", F_TITLE)
    put(ws, "A2", "Signs flipped so credits show positive. Equity is pre-closing (excludes CY income).", F_NOTE)
    signoff_row(ws)
    header_row(ws, 4, ["Financial statement line", "WP ref", "CY 12/31/2025", "PY 12/31/2024",
                       "$ change", "% change", "Notes / tickmarks"])
    lines = [("Cash", 1, "C"), ("Accounts receivable, net", 1, "WP2"), ("Inventory", 1, "D"),
             ("Prepaid expenses", 1, "E"), ("Property and equipment, net", 1, "F"), ("TOTAL_ASSETS", 0, ""),
             ("Accounts payable", -1, "WP3"), ("Accrued liabilities", -1, "WP3"),
             ("Income taxes payable", -1, "G"), ("Line of credit", -1, "H"), ("TOTAL_LIAB", 0, ""),
             ("Equity", -1, "I"), ("BLANK", 0, ""),
             ("Net sales", -1, "WP1"), ("Cost of goods sold", 1, "D"), ("Operating expenses", 1, "J"),
             ("Interest expense", 1, "H"), ("PRETAX", 0, ""), ("Income tax expense", 1, "G"),
             ("NET_INCOME", 0, ""), ("CHECK", 0, "")]
    rowmap = {}
    r = 5
    for name, sign, ref in lines:
        rowmap[name] = r
        if name == "BLANK":
            r += 1
            continue
        if sign != 0:
            put(ws, f"A{r}", name); put(ws, f"B{r}", ref)
            pre = "-" if sign < 0 else ""
            put(ws, f"C{r}", f"={pre}SUMIFS(D_TB!$E:$E,D_TB!$C:$C,$A{r})", F_LINK, MONEY0)
            put(ws, f"D{r}", f"={pre}SUMIFS(D_TB!$F:$F,D_TB!$C:$C,$A{r})", F_LINK, MONEY0)
        r += 1
    a0, a1 = rowmap["Cash"], rowmap["Property and equipment, net"]
    l0, l1 = rowmap["Accounts payable"], rowmap["Line of credit"]
    formulas = {
        "TOTAL_ASSETS": ("Total assets", "=SUM({c}%d:{c}%d)" % (a0, a1)),
        "TOTAL_LIAB": ("Total liabilities", "=SUM({c}%d:{c}%d)" % (l0, l1)),
        "PRETAX": ("Income before taxes",
                   "={c}%d-{c}%d-{c}%d-{c}%d" % (rowmap["Net sales"], rowmap["Cost of goods sold"],
                                                 rowmap["Operating expenses"], rowmap["Interest expense"])),
        "NET_INCOME": ("Net income", "={c}%d-{c}%d" % (rowmap["PRETAX"], rowmap["Income tax expense"])),
        "CHECK": ("Check: assets - liabilities - equity - net income (should be 0)",
                  "=ROUND({c}%d-{c}%d-{c}%d-{c}%d,2)" % (rowmap["TOTAL_ASSETS"], rowmap["TOTAL_LIAB"],
                                                        rowmap["Equity"], rowmap["NET_INCOME"])),
    }
    for key, (label, f) in formulas.items():
        rr = rowmap[key]
        put(ws, f"A{rr}", label, F_BOLD)
        for c in ("C", "D"):
            put(ws, f"{c}{rr}", f.format(c=c), F_BOLD, MONEY0)
    for rr in range(5, r):
        if ws[f"C{rr}"].value is not None and rr != rowmap["CHECK"]:
            put(ws, f"E{rr}", f"=C{rr}-D{rr}", fmt=MONEY0)
            put(ws, f"F{rr}", f'=IF(D{rr}=0,"",E{rr}/ABS(D{rr}))', fmt=PCT)
    ws[f"E{rowmap['CHECK']}"].value = None
    ws[f"F{rowmap['CHECK']}"].value = None
    put(ws, f"A{r + 1}", "Gross margin %", F_BOLD)
    for c in ("C", "D"):
        put(ws, f"{c}{r + 1}",
            f"=IF({c}{rowmap['Net sales']}=0,\"\",1-{c}{rowmap['Cost of goods sold']}/{c}{rowmap['Net sales']})",
            fmt=PCT)
    put(ws, f"A{r + 2}", "Gross A/R to sales (days)", F_BOLD)
    put(ws, f"C{r + 2}", f"=SUMIFS(D_TB!$E:$E,D_TB!$A:$A,\"1100\")/C{rowmap['Net sales']}*365", F_LINK, "0.0")
    put(ws, f"D{r + 2}", f"=SUMIFS(D_TB!$F:$F,D_TB!$A:$A,\"1100\")/D{rowmap['Net sales']}*365", F_LINK, "0.0")
    for ref in ("G5:G16", "G18:G24", f"G{r + 1}:G{r + 2}"):
        style_range(ws, ref, font=F_INPUT, fill=FILL_INPUT, border=BOX,
                    align=Alignment(wrap_text=True, vertical="top"))
    widths(ws, {"A": 44, "B": 8, "C": 16, "D": 16, "E": 14, "F": 10, "G": 40})
    ws.freeze_panes = "A5"
    LEAD_SALES = f"Lead!$C${rowmap['Net sales']}"

    # ================================================================ Planning
    ws = wb["Planning"]
    put(ws, "A1", "Northfield Supply Co. | FY2025 | Planning: Materiality and Risks", F_TITLE)
    put(ws, "A2", "Private company, US GAAP, AICPA AU-C standards.", F_NOTE)
    signoff_row(ws)
    section(ws, 4, "Benchmarks (from lead sheet)", 6)
    put(ws, "A5", "Net sales"); put(ws, "B5", f"={LEAD_SALES}", F_LINK, MONEY0)
    put(ws, "A6", "Income before taxes"); put(ws, "B6", f"=Lead!$C${rowmap['PRETAX']}", F_LINK, MONEY0)
    put(ws, "A7", "Total assets"); put(ws, "B7", f"=Lead!$C${rowmap['TOTAL_ASSETS']}", F_LINK, MONEY0)
    section(ws, 9, "Materiality (AU-C 320)", 6)
    put(ws, "A10", "Benchmark selected"); inp(ws, "B10", "Net sales")
    put(ws, "A11", "Benchmark amount"); put(ws, "B11", "=B5", fmt=MONEY0)
    put(ws, "C11", "Change this reference if you pick a different benchmark.", F_NOTE)
    put(ws, "A12", "Materiality % of benchmark"); inp(ws, "B12", 0.006, PCT)
    put(ws, "C12", "Common ranges: 0.5%-1% of revenue, 5% of pretax income. Justify your choice below.", F_NOTE)
    put(ws, "A13", "Overall materiality (OM)", F_BOLD); put(ws, "B13", "=ROUND(B11*B12,-3)", F_BOLD, MONEY0)
    put(ws, "A14", "Performance materiality %"); inp(ws, "B14", 0.75, PCT)
    put(ws, "C14", "50%-75% of OM depending on expected misstatements and control environment.", F_NOTE)
    put(ws, "A15", "Performance materiality (PM)", F_BOLD); put(ws, "B15", "=ROUND(B13*B14,-3)", F_BOLD, MONEY0)
    put(ws, "A16", "Clearly trivial %"); inp(ws, "B16", 0.05, PCT)
    put(ws, "A17", "Clearly trivial threshold", F_BOLD); put(ws, "B17", "=ROUND(B13*B16,-2)", F_BOLD, MONEY0)
    put(ws, "A18", "Rationale for benchmark"); inp(ws, "B18")
    ws.merge_cells("B18:F18")
    section(ws, 20, "Significant risks and planned responses (AU-C 315 / 240 / 330)", 6)
    header_row(ws, 21, ["Risk", "Assertion(s)", "Account", "Fraud risk?", "Planned response", "WP ref"])
    example = ["EXAMPLE: Improper revenue recognition (presumed fraud risk)", "Occurrence, cutoff",
               "Net sales / A/R", "Yes", "Monthly analytics, cutoff test, duplicate test, receipts testing",
               "WP1, WP2"]
    for j, v in enumerate(example, start=1):
        c = ws.cell(row=22, column=j, value=v); c.font = F_NOTE; c.fill = FILL_EXAMPLE
    for rr in range(23, 28):
        style_range(ws, f"A{rr}:F{rr}", font=F_INPUT, fill=FILL_INPUT, border=BOX)
    widths(ws, {"A": 34, "B": 22, "C": 18, "D": 12, "E": 50, "F": 10})
    PM, OM, TRIV = "Planning!$B$15", "Planning!$B$13", "Planning!$B$17"

    # ================================================================ WP1 Revenue
    ws = wb["WP1_Revenue"]
    wp_header(ws, "Revenue: substantive analytics, cutoff, and duplicate testing", "WP1")
    put(ws, "A5", "Objective: obtain evidence that net sales occurred, are complete, and are recorded in the "
                  "correct period.", F_NOTE)
    section(ws, 7, "A. Monthly substantive analytics (AU-C 520)", 10)
    put(ws, "A8", "Growth assumption"); inp(ws, "B8", 0.065, PCT)
    put(ws, "C8", "Source: management inquiry (6-7% volume growth). Corroborate before relying on it.", F_NOTE)
    put(ws, "D8", "Basis and corroborating evidence", F_BOLD)
    style_range(ws, "E8:J9", font=F_INPUT, fill=FILL_INPUT, border=BOX,
                align=Alignment(wrap_text=True, vertical="top"))
    ws.merge_cells("E8:J9")
    put(ws, "A9", "Investigation threshold"); put(ws, "B9", f"=ROUND({PM}*0.5,-3)", F_LINK, MONEY0)
    put(ws, "C9", "Default = 50% of PM per month. Override in yellow if you justify another level.", F_NOTE)
    header_row(ws, 10, ["Month", "PY actual", "Expected CY", "CY recorded", "Difference", "% diff",
                        "Over threshold?", "Explanation and corroborating evidence", "Tickmark"])
    for i in range(12):
        rr = 11 + i
        put(ws, f"A{rr}", f"2025-{i + 1:02d}")
        put(ws, f"B{rr}", f'=SUMIFS(D_PYMonthly!$B:$B,D_PYMonthly!$A:$A,"2024-"&RIGHT(A{rr},2))', F_LINK, MONEY0)
        put(ws, f"C{rr}", f"=B{rr}*(1+$B$8)", fmt=MONEY0)
        put(ws, f"D{rr}", f"=SUMIFS(D_Sales!$F:$F,D_Sales!$H:$H,A{rr})", F_LINK, MONEY0)
        put(ws, f"E{rr}", f"=D{rr}-C{rr}", fmt=MONEY0)
        put(ws, f"F{rr}", f'=IF(C{rr}=0,"",E{rr}/C{rr})', fmt=PCT)
        put(ws, f"G{rr}", f'=IF(ABS(E{rr})>$B$9,"INVESTIGATE","")', F_BOLD)
        style_range(ws, f"H{rr}:I{rr}", font=F_INPUT, fill=FILL_INPUT, border=BOX)
    put(ws, "A23", "Total", F_BOLD)
    for c in "BCDE":
        put(ws, f"{c}23", f"=SUM({c}11:{c}22)", F_BOLD, MONEY0)
    put(ws, "A24", "Difference to lead sheet (should be 0)"); put(ws, "D24", f"=ROUND(D23-{LEAD_SALES},2)", F_LINK, MONEY)

    section(ws, 26, "B. Cutoff test: late-December invoices (FOB shipping point)", 10)
    put(ws, "A27", "Suggested selection: invoices dated 12/23-12/31 of $15,000 or more. The January 2026 "
                   "sales register was not provided, so completeness cutoff is outside this exercise.", F_NOTE)
    put(ws, "E27", "Selection method / rationale", F_BOLD)
    style_range(ws, "F27:J27", font=F_INPUT, fill=FILL_INPUT, border=BOX,
                align=Alignment(wrap_text=True, vertical="top"))
    ws.merge_cells("F27:J27")
    header_row(ws, 28, ["Invoice", "Invoice date", "Ship date (ledger)", "Ship date (document)",
                        "Customer", "Amount", "Shipping doc agreed?", "Correct period? (Y/N)",
                        "Misstatement", "Notes / tickmark"])
    cut = sales[(sales["inv_date"] >= "2025-12-23") & (sales["amount"] >= 15000)].sort_values("inv_date")
    ship_by_inv = ship.drop_duplicates("inv_no").set_index("inv_no")
    rr = 29
    for _, s in cut.iterrows():
        put(ws, f"A{rr}", s["inv_no"])
        put(ws, f"B{rr}", s["inv_date"].to_pydatetime(), fmt=DATE)
        put(ws, f"C{rr}", None if pd.isna(s["ship_date"]) else s["ship_date"].to_pydatetime(), fmt=DATE)
        doc_date = ship_by_inv.loc[s["inv_no"], "ship_date"] if s["inv_no"] in ship_by_inv.index else pd.NaT
        put(ws, f"D{rr}", None if pd.isna(doc_date) else doc_date.to_pydatetime(), F_LINK, DATE)
        put(ws, f"E{rr}", s["customer"])
        put(ws, f"F{rr}", float(s["amount"]), fmt=MONEY)
        style_range(ws, f"G{rr}:H{rr}", font=F_INPUT, fill=FILL_INPUT, border=BOX)
        inp(ws, f"I{rr}", None, MONEY)
        ws[f"I{rr}"].border = BOX
        inp(ws, f"J{rr}")
        rr += 1
    put(ws, f"A{rr}", "Total selected / misstatement", F_BOLD)
    put(ws, f"F{rr}", f"=SUM(F29:F{rr - 1})", F_BOLD, MONEY)
    put(ws, f"I{rr}", f"=SUM(I29:I{rr - 1})", F_BOLD, MONEY)
    put(ws, f"A{rr + 1}", "Coverage of December revenue")
    put(ws, f"F{rr + 1}", f"=F{rr}/D22", fmt=PCT)
    add_list_validation(ws, f"G29:G{rr - 1}", ["Yes", "No"])
    add_list_validation(ws, f"H29:H{rr - 1}", ["Y", "N"])

    dstart = rr + 3
    section(ws, dstart, "C. Duplicate invoice test (full population)", 10)
    put(ws, f"A{dstart + 1}", "Test D_Sales for repeated invoice numbers, and for same customer + amount + date. "
                              "List what you find and your disposition.", F_NOTE)
    header_row(ws, dstart + 2, ["Invoice", "Times posted", "Customer", "Amount", "Disposition",
                                "Misstatement", "Notes"])
    for k in range(dstart + 3, dstart + 8):
        style_range(ws, f"A{k}:G{k}", font=F_INPUT, fill=FILL_INPUT, border=BOX)
    cstart = dstart + 9
    section(ws, cstart, "Conclusion", 10)
    inp(ws, f"A{cstart + 1}")
    ws.merge_cells(f"A{cstart + 1}:J{cstart + 3}")
    ws[f"A{cstart + 1}"].alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 22, "B": 14, "C": 16, "D": 18, "E": 34, "F": 14, "G": 16,
                "H": 16, "I": 16, "J": 44})
    ws.freeze_panes = "A5"

    # ================================================================ WP2 AR
    ws = wb["WP2_AR"]
    wp_header(ws, "Accounts receivable: existence and valuation", "WP2")
    put(ws, "A5", "Objective: A/R exists and is stated at net realizable value. Confirmations (AU-C 505) are "
                  "simulated by subsequent receipts testing in this exercise.", F_NOTE)
    section(ws, 7, "A. Population and tie-out", 8)
    items = [("Aging detail total", "=SUM(D_AR!$F:$F)", F_LINK),
             ("0-30", "=SUM(D_AR!$G:$G)", F_LINK), ("31-60", "=SUM(D_AR!$H:$H)", F_LINK),
             ("61-90", "=SUM(D_AR!$I:$I)", F_LINK), ("91-120", "=SUM(D_AR!$J:$J)", F_LINK),
             ("Over 120", "=SUM(D_AR!$K:$K)", F_LINK),
             ("Cross-foot difference (should be 0)", "=ROUND(B8-SUM(B9:B13),2)", F_BASE),
             ("TB 1100 A/R", '=SUMIFS(D_TB!$E:$E,D_TB!$A:$A,"1100")', F_LINK),
             ("Difference to TB (should be 0)", "=ROUND(B8-B15,2)", F_BASE),
             ("Balances over 90 days", "=B12+B13", F_BASE),
             ("  % of total", '=IF(B8=0,"",B17/B8)', F_BASE),
             ("Balances on credit hold", '=SUMIFS(D_AR!$F:$F,D_AR!$L:$L,"Y")', F_LINK),
             ("Customers not in master (UNMATCHED)", '=SUMIFS(D_AR!$F:$F,D_AR!$A:$A,"UNMATCHED")', F_LINK),
             ("Open items (count)", "=COUNTA(D_AR!$C:$C)-1", F_LINK)]
    for i, (label, f, font) in enumerate(items, start=8):
        put(ws, f"A{i}", label)
        put(ws, f"B{i}", f, font, PCT if "%" in label else ("0" if "count" in label else MONEY))

    section(ws, 23, "B. Sample selection and subsequent receipts (AU-C 530 / 500)", 11)
    put(ws, "A24", "Enter invoice numbers and selection basis. For 'see remit' receipts, match by customer "
                   "and amount. Include every item >= PM, then select the remainder using the documented method.",
        F_NOTE)
    ws.merge_cells("A24:E24")
    ws["A24"].alignment = Alignment(wrap_text=True, vertical="top")
    put(ws, "F24", "Sample design / rationale", F_BOLD)
    style_range(ws, "G24:K24", font=F_INPUT, fill=FILL_INPUT, border=BOX,
                align=Alignment(wrap_text=True, vertical="top"))
    ws.merge_cells("G24:K24")
    header_row(ws, 25, ["#", "Selection basis", "Invoice", "Customer", "Balance 12/31", "Receipt date",
                        "Receipt amount", "Unreceived / exception", "Alternative procedure", "Tickmark", "Notes"])
    ex = aging.sort_values("balance", ascending=False).iloc[9]
    for k in range(26, 51):
        rr = k
        n = k - 25
        put(ws, f"A{rr}", "EX" if n == 1 else n, F_NOTE if n == 1 else F_BASE)
        if n == 1:
            put(ws, f"B{rr}", "EXAMPLE: key item", F_NOTE, fill=FILL_EXAMPLE)
            put(ws, f"C{rr}", ex["inv_no"], F_NOTE, fill=FILL_EXAMPLE)
        else:
            inp(ws, f"B{rr}"); inp(ws, f"C{rr}")
        put(ws, f"D{rr}", f'=IF(C{rr}="","",IFERROR(INDEX(D_AR!$B:$B,MATCH(C{rr},D_AR!$C:$C,0)),"NOT IN AGING"))', F_LINK)
        put(ws, f"E{rr}", f'=IF(C{rr}="","",IFERROR(INDEX(D_AR!$F:$F,MATCH(C{rr},D_AR!$C:$C,0)),0))', F_LINK, MONEY)
        put(ws, f"F{rr}", f'=IF(C{rr}="","",IFERROR(INDEX(D_Receipts!$A:$A,MATCH(C{rr},D_Receipts!$D:$D,0)),"none"))', F_LINK, DATE)
        put(ws, f"G{rr}", f'=IF(C{rr}="","",SUMIFS(D_Receipts!$E:$E,D_Receipts!$D:$D,C{rr}))', F_LINK, MONEY)
        put(ws, f"H{rr}", f'=IF(C{rr}="","",E{rr}-G{rr})', fmt=MONEY)
        style_range(ws, f"I{rr}:K{rr}", font=F_INPUT, fill=FILL_INPUT, border=BOX)
        if n == 1:
            style_range(ws, f"I{rr}:K{rr}", font=F_NOTE, fill=FILL_EXAMPLE)
    put(ws, "A51", "Totals", F_BOLD)
    put(ws, "E51", "=SUM(E27:E50)", F_BOLD, MONEY)
    put(ws, "G51", "=SUM(G27:G50)", F_BOLD, MONEY)
    put(ws, "H51", "=SUM(H27:H50)", F_BOLD, MONEY)
    put(ws, "A52", "Coverage of A/R (excl. example row)"); put(ws, "E52", '=IF(B8=0,"",E51/B8)', fmt=PCT)

    section(ws, 54, "C. Allowance for doubtful accounts (AU-C 540)", 8)
    put(ws, "A55", "Gross A/R"); put(ws, "B55", "=B8", fmt=MONEY)
    put(ws, "A56", "Recorded allowance (TB 1150)"); put(ws, "B56", '=-SUMIFS(D_TB!$E:$E,D_TB!$A:$A,"1150")', F_LINK, MONEY)
    put(ws, "A57", "Recorded allowance % of gross"); put(ws, "B57", '=IF(B55=0,"",B56/B55)', fmt=PCT)
    header_row(ws, 59, ["Specific reserve: customer (exact name from D_AR)", "Balance", "Reserve %",
                        "Specific reserve", "Basis / evidence"])
    for k in range(60, 65):
        inp(ws, f"A{k}")
        put(ws, f"B{k}", f'=IF(A{k}="",0,SUMIFS(D_AR!$F:$F,D_AR!$B:$B,A{k}))', F_LINK, MONEY)
        inp(ws, f"C{k}", None, PCT)
        put(ws, f"D{k}", f"=B{k}*C{k}", fmt=MONEY)
        inp(ws, f"E{k}")
    put(ws, "A66", "General reserve %"); inp(ws, "B66", 0.01, PCT)
    put(ws, "D66", "General reserve basis / evidence", F_BOLD)
    style_range(ws, "E66:K66", font=F_INPUT, fill=FILL_INPUT, border=BOX,
                align=Alignment(wrap_text=True, vertical="top"))
    ws.merge_cells("E66:K66")
    put(ws, "A67", "Balance subject to general reserve"); put(ws, "B67", "=B55-SUM(B60:B64)", fmt=MONEY)
    put(ws, "A68", "General reserve"); put(ws, "B68", "=B67*B66", fmt=MONEY)
    put(ws, "A69", "Required allowance", F_BOLD); put(ws, "B69", "=SUM(D60:D64)+B68", F_BOLD, MONEY)
    put(ws, "A70", "Recorded allowance"); put(ws, "B70", "=B56", fmt=MONEY)
    put(ws, "A71", "Difference (positive = allowance understated)", F_BOLD)
    put(ws, "B71", "=ROUND(B69-B70,2)", F_BOLD, MONEY)
    put(ws, "A72", "Compare to clearly trivial / PM"); put(ws, "B72", f"={TRIV}", F_LINK, MONEY0)
    put(ws, "C72", f"={PM}", F_LINK, MONEY0)
    section(ws, 74, "Conclusion", 8)
    inp(ws, "A75"); ws.merge_cells("A75:K77")
    ws["A75"].alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 40, "B": 18, "C": 14, "D": 32, "E": 16, "F": 13, "G": 15, "H": 16, "I": 26, "J": 10, "K": 30})
    ws.freeze_panes = "A5"

    # ================================================================ WP3 SURL
    ws = wb["WP3_SURL"]
    wp_header(ws, "Search for unrecorded liabilities: January 2026 disbursements", "WP3")
    put(ws, "A5", "Objective: A/P and accrued liabilities are complete at 12/31/2025. For each payment, decide "
                  "whether the goods or services were received on or before 12/31/2025.", F_NOTE)
    put(ws, "A7", "Testing threshold"); put(ws, "B7", f"={TRIV}", F_LINK, MONEY0)
    put(ws, "C7", "Default = clearly trivial. Test every payment at or above it; note smaller items you spot.", F_NOTE)
    header_row(ws, 9, ["Check/ACH", "Paid", "Vendor", "Vendor invoice", "Invoice date", "Amount", "Description",
                       ">= threshold?", "On 12/31 A/P listing?", "Goods/services period", "12/31 liability? (Y/N)",
                       "Unrecorded amount", "Notes / tickmark"])
    rr = 10
    for _, d in jd.sort_values("paid_date").iterrows():
        put(ws, f"A{rr}", d["check_no"])
        put(ws, f"B{rr}", d["paid_date"].to_pydatetime(), fmt=DATE)
        put(ws, f"C{rr}", d["vendor"])
        put(ws, f"D{rr}", d["vendor_inv"])
        put(ws, f"E{rr}", d["inv_date"].to_pydatetime(), fmt=DATE)
        put(ws, f"F{rr}", float(d["amount"]), fmt=MONEY)
        put(ws, f"G{rr}", d["description"])
        put(ws, f"H{rr}", f'=IF(F{rr}>=$B$7,"Test","")')
        put(ws, f"I{rr}", f'=IF(COUNTIFS(D_AP!$B:$B,D{rr})>0,"Yes","No")', F_LINK)
        style_range(ws, f"J{rr}:K{rr}", font=F_INPUT, fill=FILL_INPUT, border=BOX)
        put(ws, f"L{rr}", f'=IF(AND(UPPER(K{rr})="Y",I{rr}="No"),F{rr},0)', fmt=MONEY)
        inp(ws, f"M{rr}")
        rr += 1
    add_list_validation(ws, f"K10:K{rr - 1}", ["Y", "N"])
    put(ws, f"A{rr}", "Totals", F_BOLD)
    put(ws, f"F{rr}", f"=SUM(F10:F{rr - 1})", F_BOLD, MONEY)
    put(ws, f"L{rr}", f"=SUM(L10:L{rr - 1})", F_BOLD, MONEY)
    put(ws, f"A{rr + 1}", "Tie to D_JanDisb (should be 0)")
    put(ws, f"F{rr + 1}", f"=ROUND(F{rr}-SUM(D_JanDisb!$F:$F),2)", F_LINK, MONEY)
    put(ws, f"A{rr + 2}", "Unrecorded total vs PM"); put(ws, f"F{rr + 2}", f"={PM}", F_LINK, MONEY0)
    put(ws, f"L{rr + 2}", f'=IF(L{rr}>F{rr + 2},"EXCEEDS PM","Below PM")', F_BOLD)
    section(ws, rr + 4, "Conclusion", 13)
    inp(ws, f"A{rr + 5}"); ws.merge_cells(f"A{rr + 5}:M{rr + 7}")
    ws[f"A{rr + 5}"].alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 12, "B": 12, "C": 28, "D": 15, "E": 12, "F": 14, "G": 42, "H": 12, "I": 12,
                "J": 20, "K": 12, "L": 15, "M": 30})
    ws.freeze_panes = "A10"

    # ================================================================ Findings
    ws = wb["Findings"]
    put(ws, "A1", "Findings and summary of misstatements (AU-C 450)", F_TITLE)
    put(ws, "A2", "One row per finding. 'Reference' must hold the invoice #, vendor invoice #, check #, or "
                  "customer name so score_results.py can match it. Amount = overstatement of income/assets "
                  "or understatement of liabilities, as a positive number.", F_NOTE)
    signoff_row(ws)
    hdrs = ["ID", "WP ref", "Error type", "Reference", "Amount", "Accounts affected",
            "Factual / judgmental / projected", "Caught by", "Found at (hh:mm)", "Notes"]
    header_row(ws, 4, hdrs)
    exrow = ["EX-1", "WP1", "EXAMPLE: cutoff", "INV-XXXXXX", 12345.67, "Sales / A/R", "Factual",
             "AI-assisted", "02:40", "Delete nothing; scoring skips IDs starting with EX"]
    for j, v in enumerate(exrow, start=1):
        c = ws.cell(row=5, column=j, value=v); c.font = F_NOTE; c.fill = FILL_EXAMPLE
    ws["E5"].number_format = MONEY
    for k in range(6, 26):
        style_range(ws, f"A{k}:J{k}", font=F_INPUT, fill=FILL_INPUT, border=BOX)
        ws[f"E{k}"].number_format = MONEY
    add_list_validation(ws, "G6:G25", ["Factual", "Judgmental", "Projected"])
    add_list_validation(ws, "H6:H25", ["AI-assisted", "Human review", "Reviewer agent"])
    put(ws, "H3", "Caught by: AI-assisted | Human review | Reviewer agent", F_NOTE)
    put(ws, "D27", "Gross identified", F_BOLD); put(ws, "E27", "=SUM(E6:E25)", F_BOLD, MONEY)
    put(ws, "F27", "Confirm related findings do not double-count the same proposed adjustment.", F_NOTE)
    put(ws, "D28", "Overall materiality"); put(ws, "E28", f"={OM}", F_LINK, MONEY0)
    put(ws, "D29", "Performance materiality"); put(ws, "E29", f"={PM}", F_LINK, MONEY0)
    put(ws, "D30", "Aggregate vs OM", F_BOLD)
    put(ws, "E30", '=IF(E27>=E28,"MATERIAL - discuss with partner",IF(E27>=E29,"Above PM - evaluate","Below PM"))', F_BOLD)
    section(ws, 32, "Evaluation and next step", 10)
    style_range(ws, "A33:J35", font=F_INPUT, fill=FILL_INPUT, border=BOX,
                align=Alignment(wrap_text=True, vertical="top"))
    ws.merge_cells("A33:J35")
    widths(ws, {"A": 8, "B": 8, "C": 30, "D": 24, "E": 16, "F": 22, "G": 16, "H": 16, "I": 12, "J": 44})
    ws.freeze_panes = "A5"

    # ================================================================ Review notes, AI log, time log
    def log_sheet(name, title, headers, example, n=20, w=None):
        s = wb[name]
        put(s, "A1", title, F_TITLE)
        header_row(s, 3, headers)
        for j, v in enumerate(example, start=1):
            c = s.cell(row=4, column=j, value=v); c.font = F_NOTE; c.fill = FILL_EXAMPLE
        for k in range(5, 5 + n):
            style_range(s, f"A{k}:{get_column_letter(len(headers))}{k}", font=F_INPUT, fill=FILL_INPUT, border=BOX)
        for i, width in enumerate(w or [], start=1):
            s.column_dimensions[get_column_letter(i)].width = width
        s.freeze_panes = "A4"
        return s

    s = log_sheet("Review_Notes", "Review notes (AU-C 230)",
                  ["ID", "WP ref", "Cell / area", "Comment", "Raised by", "Severity", "Preparer response", "Status"],
                  ["EX", "WP2", "B71", "EXAMPLE: Support for the 1% general reserve rate is not documented.",
                   "Reviewer agent", "Medium", "Added 3-year write-off history.", "Cleared"],
                  w=[6, 8, 12, 60, 16, 10, 50, 10])
    add_list_validation(s, "F5:F24", ["High", "Medium", "Low"])
    add_list_validation(s, "H5:H24", ["Open", "Cleared"])
    log_sheet("AI_Log", "AI failure log and product feedback",
              ["#", "Block", "What I asked the AI to do", "What it got wrong", "How I caught it",
               "Severity", "Product suggestion"],
              ["EX", "WP1", "EXAMPLE: Sum monthly revenue from the raw ledger",
               "Included subtotal rows, doubling revenue", "Total did not agree to TB",
               "High", "Auto-detect and drop subtotal rows on import, with a visible log"],
              w=[5, 10, 40, 40, 32, 10, 44])
    s = log_sheet("Time_Log", "Time log (minutes)",
                  ["Block", "Planned", "Actual", "Variance", "What AI did", "Notes"],
                  ["EXAMPLE", 30, 35, "", "Drafted memo", ""], n=0, w=[36, 10, 10, 10, 44, 40])
    blocks = [("1 Run pipeline + inspect data", 60), ("2 Planning memo", 30), ("3 Lead sheet review", 30),
              ("4 WP1 Revenue", 75), ("5 WP2 A/R", 75), ("Break", 15), ("6 WP3 SURL", 60),
              ("7 Reviewer agent + clear notes", 60), ("8 Results + failure log", 45), ("9 Package + Loom", 30)]
    for i, (b, m) in enumerate(blocks, start=5):
        put(s, f"A{i}", b); put(s, f"B{i}", m)
        inp(s, f"C{i}"); put(s, f"D{i}", f'=IF(C{i}="","",C{i}-B{i})'); inp(s, f"E{i}"); inp(s, f"F{i}")
    t = 5 + len(blocks)
    put(s, f"A{t}", "Total", F_BOLD)
    for c in "BC":
        put(s, f"{c}{t}", f"=SUM({c}5:{c}{t - 1})", F_BOLD)
    put(s, f"D{t}", f"=C{t}-B{t}", F_BOLD)

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = sheet.title.startswith("D_")
    wb.save(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
