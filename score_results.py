#!/usr/bin/env python3
"""
score_results.py
Compares the Findings sheet in your workbook to answer_key/answer_key.json.
Run this ONLY in the Results block, after your findings are final.

Usage
  python score_results.py --wb Northfield_FY2025_Workpapers.xlsx --key answer_key/answer_key.json
Writes results.md next to the workbook.
"""
import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def load_findings(path):
    ws = load_workbook(path, data_only=True)["Findings"]
    out = []
    for row in ws.iter_rows(min_row=5, max_row=25, max_col=10, values_only=True):
        fid = row[0]
        if fid is None or str(fid).upper().startswith("EX"):
            continue
        out.append({"id": str(fid), "wp": row[1], "type": str(row[2] or ""), "ref": str(row[3] or ""),
                    "amount": float(row[4] or 0), "caught_by": str(row[7] or "Unspecified"),
                    "notes": str(row[9] or "")})
    return out


def matches(item, f):
    blob = " ".join([f["ref"], f["type"], f["notes"]]).upper()
    if any(r.upper() in blob for r in item.get("match_refs", [])):
        return True
    keywords = item.get("match_keywords", [])
    keyword_hits = sum(1 for k in keywords if k.upper() in blob)
    return keyword_hits >= item.get("min_keyword_hits", 1) if keywords else False


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--wb", default="Northfield_FY2025_Workpapers.xlsx")
    ap.add_argument("--key", default="answer_key/answer_key.json")
    args = ap.parse_args()
    key = json.loads(Path(args.key).read_text())
    findings = load_findings(args.wb)

    used = set()
    lines = ["# Results vs answer key", "",
             f"Seed {key['seed']}. Findings logged: {len(findings)}.", "",
             "| Error | Type | Planted amount | Found? | Credit | Your amount | Caught by |",
             "|---|---|---:|---|---:|---:|---|"]
    detection_points = 0.0
    by_method = {}
    for item in key["errors"]:
        hits = [i for i, f in enumerate(findings) if i not in used and matches(item, f)]
        used.update(hits)
        if hits:
            amt = sum(findings[i]["amount"] for i in hits)
            methods = sorted(set(findings[i]["caught_by"] for i in hits))
            method = ", ".join(methods)
            inv_refs = [r for r in item.get("match_refs", []) if r.upper().startswith("INV-")]
            status = "Yes"
            credit = 1.0
            if len(inv_refs) > 1:
                blob = " ".join(findings[i]["ref"] + " " + findings[i]["notes"] for i in hits).upper()
                got = sum(1 for r in inv_refs if r.upper() in blob)
                status = "Yes" if got == len(inv_refs) else f"Partial ({got}/{len(inv_refs)} invoices)"
                credit = got / len(inv_refs)
            detection_points += credit
            for m in methods:
                by_method[m] = by_method.get(m, 0) + credit / len(methods)
            lines.append(f"| {item['id']} | {item['type']} | {item['amount']:,.2f} | {status} | "
                         f"{credit:.2f} | {amt:,.2f} | {method} |")
        else:
            lines.append(f"| {item['id']} | {item['type']} | {item['amount']:,.2f} | **Missed** | 0.00 | - | - |")

    bonus_hits = [b["id"] for b in key.get("bonus_not_scored", [])
                  if any(matches(b, f) for f in findings)]
    for b in key.get("bonus_not_scored", []):
        used.update(i for i, f in enumerate(findings) if matches(b, f))
    unmatched = [findings[i] for i in range(len(findings)) if i not in used]

    total = len(key["errors"])
    lines += ["", f"**Detection score: {detection_points:.2f}/{total} ({detection_points / total:.0%})**", ""]
    lines.append("Caught by: " + (", ".join(f"{k} {v:.2f}" for k, v in by_method.items()) or "n/a"))
    lines.append(f"Bonus (not scored): {', '.join(bonus_hits) if bonus_hits else 'none found'}")
    lines += ["", "## Findings that match no planted error",
              "These are either false positives or legitimate extra observations. Review each one.", ""]
    lines += [f"- {f['id']} ({f['wp']}): {f['type']} | {f['ref']} | {f['amount']:,.2f}" for f in unmatched] or ["- none"]
    lines += ["", "## Answer key detail", ""]
    for item in key["errors"]:
        lines.append(f"- **{item['id']} {item['type']}**: {item['description']}")
    lines += ["", "Decoys: " + " ".join(key.get("decoys", []))]

    out = Path(args.wb).with_name("results.md")
    out.write_text("\n".join(lines) + "\n")
    print("\n".join(lines[:6 + total + 5]))
    print(f"\nFull report: {out}")


if __name__ == "__main__":
    main()
